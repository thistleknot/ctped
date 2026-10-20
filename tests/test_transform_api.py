"""Stage | Mechanism | Tag | Requirements
-------|-----------|-----|-------------
WORKBOOK | Build valid and invalid XLSX rule inputs | [mandatory] | R2,R5
PREVIEW | Exercise request validation and nonmutating ordered transforms | [mandatory] | R3-R9
PUBLISH | Verify overlay output and source preservation | [mandatory] | R9-R11
SERVICE | Check health, workbook template, and terminal Entry progress | [mandatory] | R1-R2,R13

R1. WHEN the health endpoint is called, return readiness without file access.
R2. WHEN the template endpoint is called, return the exact Rules worksheet headers.
R3. WHEN valid rule uploads are posted, return structured per-rule results.
R4. WHEN preview runs, source files remain unchanged.
R5. WHEN malformed rules are posted, return row-specific 422 errors.
R6. WHEN rows address ordered operations, execute them deterministically.
R7. WHEN formulas resolve another declared File, use the project resolver.
R8. WHEN a later row fails, never create a partial output directory.
R9. WHEN apply succeeds, write changed data only beneath relative overlay paths.
R10. WHEN overlay output exists or is inside the source root, reject it.
R11. WHEN errors occur, source game files remain byte-identical.
R13. WHEN API rules run, emit and close a tqdm Entry bar in the server terminal.

CLOSED
- Fixtures copy only the Gamefile and two target data Files into temporary roots.
- No native editor or installed H: drive data is required for this synthetic API battery.

Preconditions: the private default.ced/Structure fixture and Python API dependencies exist.
Failures: failed assertions, unexpected API status codes, or source-byte changes fail the suite.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import shutil

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from transform_api import RULE_HEADERS, app, create_rules_workbook, parse_rules_workbook


SAMPLE_ROOT = Path(__file__).parents[1] / ".private_samples"
CLIENT = TestClient(app)


def _workbook(rows: tuple[tuple[object, ...], ...]) -> bytes:
    """Build a rule workbook from independent row values.

    Require: rows are four-cell tuples matching the API schema.
    Guarantee: return an in-memory XLSX payload.
    Maintain: no shared or installed workbook is changed.
    Assert: first row matches the public header tuple.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rules"
    sheet.append(RULE_HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _project_paths(tmp_path: Path) -> dict[str, str | Path]:
    """Create a small Gamefile project with Advance and Age target data.

    Require: the supplied private descriptor and source data fixtures exist.
    Guarantee: return paths rooted in a disposable test GameData tree.
    Maintain: all writes stay under ``tmp_path``.
    Assert: target declarations use the copied filenames.
    """
    source_data = SAMPLE_ROOT / "GameData"
    data_root = tmp_path / "GameData"
    gamefile_directory = data_root / "default" / "gamedata"
    gamefile_directory.mkdir(parents=True)
    shutil.copy2(source_data / "default" / "gamedata" / "gamefile.txt", gamefile_directory)
    shutil.copy2(source_data / "default" / "gamedata" / "Advance.txt", gamefile_directory)
    shutil.copy2(source_data / "default" / "gamedata" / "age.txt", gamefile_directory)
    return {
        "gamefile_path": gamefile_directory / "gamefile.txt",
        "template_path": SAMPLE_ROOT / "default.ced",
        "structure_root": SAMPLE_ROOT / "Structure",
        "data_root": data_root,
    }


def _multipart(paths: dict[str, str | Path], workbook_data: bytes, **extra: str) -> dict[str, object]:
    """Build form fields and upload tuple for one endpoint request.

    Require: project path values and XLSX bytes are supplied.
    Guarantee: return TestClient-compatible multipart request arguments.
    Maintain: do not alter the workbook or project paths.
    Assert: upload uses the API's workbook field and media type.
    """
    data = {
        name: str(paths[name])
        for name in ("gamefile_path", "template_path", "structure_root")
    }
    data.update(extra)
    return {
        "data": data,
        "files": {
            "workbook": (
                "rules.xlsx",
                workbook_data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    }


def test_health_and_downloaded_template_expose_the_workbook_contract() -> None:
    """Verify service readiness and the downloadable workbook schema.

    Require: API routes and optional workbook dependencies are importable.
    Guarantee: health and template responses satisfy their endpoint contracts.
    Maintain: no project paths or source files are accessed.
    Assert: the workbook contains the exact ordered Rules headers.
    """
    assert CLIENT.get("/health").json() == {"status": "ok"}
    response = CLIENT.get("/v1/transforms/template")
    assert response.status_code == 200
    assert response.headers["content-disposition"].endswith('filename="ctpedit-transform-rules.xlsx"')
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["Rules"]
        assert tuple(cell.value for cell in workbook["Rules"][1]) == RULE_HEADERS
    finally:
        workbook.close()


def test_preview_then_apply_uses_formula_references_and_preserves_sources(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Verify ordered transformations, references, overlay bytes, and isolation.

    Require: private descriptor fixtures, a writable temporary directory, and output capture.
    Guarantee: preview and apply reflect the same ordered rule results.
    Maintain: source Gamefile data bytes remain unchanged.
    Assert: only the changed Advance file is published in the overlay.
    """
    paths = _project_paths(tmp_path)
    advance_path = paths["data_root"] / "default" / "gamedata" / "Advance.txt"
    age_path = paths["data_root"] / "default" / "gamedata" / "age.txt"
    advance_before = advance_path.read_bytes()
    age_before = age_path.read_bytes()
    workbook_data = _workbook(
        (
            ("advance.ces", "Cost", "add", 10),
            (
                "advance.ces",
                "Cost",
                "formula",
                'value + ref("age.ces", "AGE_ONE", "Age")',
            ),
            ("advance.ces", "Cost", "multiply", 2),
        )
    )
    preview = CLIENT.post(
        "/v1/transforms/preview", **_multipart(paths, workbook_data)
    )
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert [rule["changed_entries"] for rule in payload["rules"]] == [107, 107, 107]
    assert payload["changed_files"] == ["default/gamedata/Advance.txt"]
    assert advance_path.read_bytes() == advance_before
    assert age_path.read_bytes() == age_before

    output = tmp_path / "overlay"
    applied = CLIENT.post(
        "/v1/transforms/apply",
        **_multipart(paths, workbook_data, output_directory=str(output)),
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["changed_files"] == ["default/gamedata/Advance.txt"]
    output_advance = output / "default" / "gamedata" / "Advance.txt"
    assert output_advance.is_file()
    assert output_advance.read_bytes() != advance_before
    assert not (output / "default" / "gamedata" / "age.txt").exists()
    assert advance_path.read_bytes() == advance_before
    assert age_path.read_bytes() == age_before
    assert "advance.ces:Cost" in capsys.readouterr().err


def test_invalid_later_rule_returns_error_without_publishing_partial_overlay(tmp_path: Path) -> None:
    """Verify a later unresolved target prevents partial publication.

    Require: temporary project fixtures and a valid rule before the bad row.
    Guarantee: apply returns a row-specific client error.
    Maintain: source bytes and output-directory absence remain unchanged.
    Assert: no overlay appears after the second rule fails.
    """
    paths = _project_paths(tmp_path)
    advance_path = paths["data_root"] / "default" / "gamedata" / "Advance.txt"
    original = advance_path.read_bytes()
    output = tmp_path / "must-not-exist"
    workbook_data = _workbook(
        (
            ("advance.ces", "Cost", "multiply", 2),
            ("missing.ces", "Cost", "multiply", 2),
        )
    )
    response = CLIENT.post(
        "/v1/transforms/apply",
        **_multipart(paths, workbook_data, output_directory=str(output)),
    )
    assert response.status_code == 422
    assert "Rules row 3" in response.json()["detail"]
    assert not output.exists()
    assert advance_path.read_bytes() == original


def test_workbook_validation_rejects_formula_cells_and_wrong_headers() -> None:
    """Verify workbook formulas and header drift are rejected.

    Require: openpyxl can construct temporary in-memory workbooks.
    Guarantee: both invalid workbook shapes raise actionable ValueErrors.
    Maintain: no project-loading or source-file operations occur.
    Assert: formula cells and altered headers are not accepted.
    """
    formula_book = Workbook()
    formula_sheet = formula_book.active
    formula_sheet.title = "Rules"
    formula_sheet.append(RULE_HEADERS)
    formula_sheet.append(("advance.ces", "Cost", "formula", "=1+2"))
    formula_bytes = BytesIO()
    formula_book.save(formula_bytes)
    formula_book.close()
    try:
        parse_rules_workbook(formula_bytes.getvalue())
    except ValueError as error:
        assert "Excel formulas are not supported" in str(error)
    else:
        raise AssertionError("Excel formula cell was accepted")

    wrong_headers = _workbook((("advance.ces", "Cost", "add", 1),))
    altered = load_workbook(BytesIO(wrong_headers))
    altered["Rules"]["B1"] = "Target"
    altered_bytes = BytesIO()
    altered.save(altered_bytes)
    altered.close()
    try:
        parse_rules_workbook(altered_bytes.getvalue())
    except ValueError as error:
        assert "headers must be" in str(error)
    else:
        raise AssertionError("workbook with wrong headers was accepted")


def test_rules_parser_accepts_each_existing_model_operation() -> None:
    """Verify workbook rows accept each existing mass-modify operation.

    Require: every row follows the four-column workbook contract.
    Guarantee: parser returns canonical operation names in worksheet order.
    Maintain: the parser does not execute or mutate project data.
    Assert: all five operations match File.mass_modify's supported set.
    """
    rules = parse_rules_workbook(
        _workbook(
            (
                ("advance.ces", "Cost", "set", 20),
                ("advance.ces", "Icon", "concatenate", "_ALT"),
                ("advance.ces", "Cost", "add", 1),
                ("advance.ces", "Cost", "multiply", 2),
                ("advance.ces", "Cost", "formula", "value * 1.1"),
            )
        )
    )
    assert [rule.operation for rule in rules] == [
        "set", "concatenate", "add", "multiply", "formula"
    ]


def test_preview_rejects_non_xlsx_upload_before_project_path_loading() -> None:
    """Verify malformed workbook bytes fail before invalid project paths load.

    Require: TestClient can post multipart form data.
    Guarantee: the API returns an actionable client error.
    Maintain: the supplied nonexistent project paths are never opened.
    Assert: response status is 422 and identifies workbook parsing.
    """
    response = CLIENT.post(
        "/v1/transforms/preview",
        data={
            "gamefile_path": "missing-gamefile.txt",
            "template_path": "missing-default.ced",
            "structure_root": "missing-Structure",
        },
        files={"workbook": ("rules.xlsx", b"not an xlsx", "application/octet-stream")},
    )
    assert response.status_code == 422
    assert "cannot read XLSX workbook" in response.json()["detail"]


def test_apply_rejects_existing_or_source_descendant_output(tmp_path: Path) -> None:
    """Verify output paths must be new and outside the source data root.

    Require: temporary project fixtures with existing and nested candidates.
    Guarantee: both invalid destinations return client errors.
    Maintain: no invalid destination is written.
    Assert: existing and source-descendant directories remain rejected.
    """
    paths = _project_paths(tmp_path)
    workbook_data = _workbook((("advance.ces", "Cost", "add", 1),))
    existing = tmp_path / "already-there"
    existing.mkdir()
    response = CLIENT.post(
        "/v1/transforms/apply",
        **_multipart(paths, workbook_data, output_directory=str(existing)),
    )
    assert response.status_code == 422
    assert "must not already exist" in response.json()["detail"]

    inside_source = paths["data_root"] / "overlay"
    response = CLIENT.post(
        "/v1/transforms/apply",
        **_multipart(paths, workbook_data, output_directory=str(inside_source)),
    )
    assert response.status_code == 422
    assert "outside the source data root" in response.json()["detail"]
    assert not inside_source.exists()
