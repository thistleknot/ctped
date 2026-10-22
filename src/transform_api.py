"""Stage | Mechanism | Tag | Requirements
-------|-----------|-----|-------------
HEALTH | Report local service readiness | [mandatory] | R1
TEMPLATE | Generate a starter Rules workbook in memory | [mandatory] | R2
INGEST | Bound and parse one uploaded XLSX rules sheet | [mandatory] | R3,R5
PROJECT | Load Gamefile, descriptor template, and explicit Structure root | [mandatory] | R4
PLAN | Apply ordered rules to fresh in-memory project models | [mandatory] | R6-R9
PUBLISH | Stage changed files as a new relative-path overlay | [mandatory] | R10-R11
PROGRESS | Render per-rule Entry traversal in the service terminal | [mandatory] | R13
SERVE | Expose preview/apply over a loopback-only FastAPI service | [mandatory] | R1,R3,R12-R13

R1. WHEN GET /health is called, return readiness without reading project data.
R2. WHEN the template route is called, return the exact Rules workbook schema.
R3. WHEN preview/apply is called, require the uploaded XLSX and project paths
    as multipart form data; apply additionally requires an output directory.
R4. WHEN a Gamefile is loaded, delegate path inference and project validation to
    app.load_gamefile with the supplied template and Structure root.
R5. WHEN a workbook is read, require one Rules sheet, exact ordered headers,
    1-500 populated rows, and <=10 MiB; reject Excel formula cells and malformed
    rule rows with their worksheet row number.
R6. WHEN a rule is executed, resolve one declared .ces and exact field path,
    apply an allowlisted operation to all Entries, and preserve worksheet order.
R7. WHEN Formula is used, delegate to File.mass_modify and
    app.resolve_project_formula_reference; formulas read the staged project at
    the beginning of their own rule and are never reevaluated.
R8. WHEN preview runs, use fresh in-memory models and return per-rule counts,
    up to ten changed-value samples, and changed relative output paths; write no
    files.
R9. WHEN apply runs, run all rules against a fresh in-memory project and stage
    only changed project data Files relative to the inferred data root.
R10. WHEN publication succeeds, rename a complete temporary sibling directory
     to the requested new output directory; reject existing, source-equal, or
     source-descendant destinations and leave source files untouched.
R11. IF workbook, project, rule, serialization, or publication validation fails,
     THEN return an actionable client error and leave no partial output folder.
R12. WHEN the service is launched from this module, bind only to 127.0.0.1.
R13. WHEN a rule runs, render an Entry-count tqdm bar in the service terminal;
     close it on success or failure and do not stream it through HTTP.

CLOSED
- The GUI remains the existing manual editing path and is not made dependent on
  FastAPI or openpyxl.
- The workbook is uploaded on every request and is never read from an arbitrary
  server-side workbook path.
- The API transforms all Entries; Entry-name filtering is outside this service.
- Apply writes a changed-file overlay, never a full game copy or source files.
- Workbook Formula operands are text in CTPEd's restricted expression grammar;
  Excel formula cells are rejected.

Preconditions: local caller supplies an accessible gamefile.txt, default.ced,
and Structure directory; declared target files are format 1/3 and editable.
Failures: invalid inputs return HTTP 422, source files remain unchanged, and
temporary staging is removed if output publication fails.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from io import BytesIO
import os
from pathlib import Path
import shutil
import tempfile
from typing import Any
from zipfile import BadZipFile

from fastapi import FastAPI, File as UploadField, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl.comments import Comment
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from tqdm import tqdm

from app import (
    LoadedProject,
    load_gamefile,
    resolve_project_formula_reference,
    resolve_project_structure_file,
)
from file import File as DataFile


MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_RULES = 500
RULE_HEADERS = ("StructureFile", "FieldPath", "Operation", "Operand")
SUPPORTED_OPERATIONS = {"set", "concatenate", "add", "multiply", "formula"}
WORKBOOK_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


@dataclass(frozen=True)
class TransformRule:
    """One ordered workbook row.

    Require: values were validated against the Rules worksheet contract.
    Guarantee: retain the source row number and canonical operation name.
    Maintain: do not mutate the project or workbook.
    Assert: operand and field identifiers are nonempty strings.
    """

    row_number: int
    structure_file: str
    field_path: str
    operation: str
    operand: str


@dataclass(frozen=True)
class RuleResult:
    """Summarize one applied in-memory rule for an API response.

    Require: samples contain only JSON-compatible scalar text.
    Guarantee: identify the rule row, changed count, and bounded examples.
    Maintain: report relative paths rather than disclose unrelated local paths.
    Assert: changed count is nonnegative.
    """

    row_number: int
    structure_file: str
    field_path: str
    operation: str
    changed_entries: int
    samples: tuple[dict[str, Any], ...]


@dataclass(frozen=True)
class TransformRun:
    """Hold completed in-memory results and changed output bytes.

    Require: every workbook rule completed without writing source files.
    Guarantee: retain ordered rule results and only byte-changed Files.
    Maintain: changed bytes are keyed by safe paths relative to the data root.
    Assert: no relative path escapes its project root.
    """

    results: tuple[RuleResult, ...]
    changed_files: tuple[tuple[Path, bytes], ...]
    data_root: Path


def create_rules_workbook() -> bytes:
    """Return the blank starter workbook.

    Require: openpyxl can construct an in-memory XLSX workbook.
    Guarantee: return one Rules worksheet with the exact API headers.
    Maintain: no workbook is written to disk.
    Assert: the produced bytes can be loaded by openpyxl.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rules"
    sheet.append(RULE_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:D1"
    for column, width in zip("ABCD", (22, 28, 18, 48)):
        sheet.column_dimensions[column].width = width
    for cell, description in zip(
        sheet[1],
        (
            "Unique declared .ces filename, such as advance.ces.",
            "Exact displayed scalar path, such as Cost or TransformAdd.Time.",
            "set, concatenate, add, multiply, or formula.",
            "Literal operand; for formula, CTPEd expression text such as value * 1.1.",
        ),
    ):
        cell.comment = Comment(description, "CTPEd Python")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse_rules_workbook(data: bytes) -> tuple[TransformRule, ...]:
    """Validate an uploaded XLSX and return its ordered rules.

    Require: data is the uploaded workbook and does not exceed the API limit.
    Guarantee: return validated rules or raise ValueError with row context.
    Maintain: do not evaluate Excel formulas or modify workbook bytes.
    Assert: all returned operation names are in SUPPORTED_OPERATIONS.
    """
    if not isinstance(data, bytes):
        raise TypeError("workbook upload must be bytes")
    if not data or len(data) > MAX_WORKBOOK_BYTES:
        raise ValueError("workbook must be nonempty and no larger than 10 MiB")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, EOFError) as error:
        raise ValueError(f"cannot read XLSX workbook: {error}") from error
    try:
        if workbook.sheetnames != ["Rules"]:
            raise ValueError("workbook must contain exactly one worksheet named Rules")
        sheet = workbook["Rules"]
        if sheet.max_column != len(RULE_HEADERS):
            raise ValueError("Rules worksheet must have exactly four columns")
        if sheet.max_row > MAX_RULES + 1:
            raise ValueError(f"Rules worksheet may contain at most {MAX_RULES} rules")
        header = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if header != RULE_HEADERS:
            raise ValueError(
                "Rules row 1 headers must be: " + ", ".join(RULE_HEADERS)
            )

        parsed: list[TransformRule] = []
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(RULE_HEADERS)), start=2
        ):
            values = tuple(cell.value for cell in cells)
            if all(value is None or value == "" for value in values):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"Rules row {row_number}: Excel formulas are not supported")
            if any(value is None or (isinstance(value, str) and not value.strip()) for value in values):
                raise ValueError(f"Rules row {row_number}: all four cells are required")
            structure_file, field_path, operation, operand_value = values
            if not isinstance(structure_file, str) or Path(structure_file).name != structure_file:
                raise ValueError(f"Rules row {row_number}: StructureFile must be a .ces basename")
            if Path(structure_file).suffix.casefold() != ".ces":
                raise ValueError(f"Rules row {row_number}: StructureFile must end in .ces")
            if not isinstance(field_path, str):
                raise ValueError(f"Rules row {row_number}: FieldPath must be text")
            if not isinstance(operation, str):
                raise ValueError(f"Rules row {row_number}: Operation must be text")
            canonical_operation = operation.strip().casefold()
            if canonical_operation not in SUPPORTED_OPERATIONS:
                raise ValueError(
                    f"Rules row {row_number}: unsupported Operation {operation!r}"
                )
            if canonical_operation == "formula" and not isinstance(operand_value, str):
                raise ValueError(f"Rules row {row_number}: Formula Operand must be text")
            if isinstance(operand_value, bool) or not isinstance(
                operand_value, (str, int, float)
            ):
                raise ValueError(f"Rules row {row_number}: Operand must be text or a number")
            operand = str(operand_value).strip() if isinstance(operand_value, str) else str(operand_value)
            if not operand:
                raise ValueError(f"Rules row {row_number}: Operand must not be blank")
            parsed.append(
                TransformRule(
                    row_number,
                    structure_file.strip(),
                    field_path.strip(),
                    canonical_operation,
                    operand,
                )
            )
        if not parsed:
            raise ValueError("Rules worksheet must contain at least one rule")
        return tuple(parsed)
    finally:
        workbook.close()


def _load_document(
    gamefile_path: str,
    template_path: str,
    structure_root: str,
    language: str | None,
) -> LoadedProject:
    """Load a fresh Gamefile-backed project for one request.

    Require: caller-supplied paths identify the Gamefile, descriptor template,
    and Structure root.
    Guarantee: return a fresh project model without importing data Files.
    Maintain: use the app's root/language inference and parser contract.
    Assert: returned roots are absolute paths.
    """
    if not gamefile_path or not template_path or not structure_root:
        raise ValueError("gamefile_path, template_path, and structure_root are required")
    return load_gamefile(gamefile_path, template_path, structure_root, language or None)


def _rule_file(document: LoadedProject, rule: TransformRule) -> DataFile:
    """Resolve one unique declared project File for a rule.

    Require: rule names a validated .ces basename.
    Guarantee: import or reuse the matching declared File.
    Maintain: follow the existing project resolver and keep Files in memory.
    Assert: the resolved File belongs to exactly one declaration.
    """
    file_model = resolve_project_structure_file(
        document, rule.structure_file, case_sensitive=False
    )
    if file_model is None:
        raise ValueError(
            f"Rules row {rule.row_number}: project has no structure {rule.structure_file!r}"
        )
    if file_model.structure.format not in {1, 3}:
        raise ValueError(
            f"Rules row {rule.row_number}: mass transformations require format 1 or 3"
        )
    if not file_model.editable:
        raise PermissionError(f"Rules row {rule.row_number}: target File is not editable")
    declared = file_model.fields.item_no_error(rule.field_path)
    observed = any(
        path.casefold() == rule.field_path.casefold()
        and not instance.field.inferred
        for entry in file_model.entries
        for path, instance, _depth in entry.field_instances_with_paths()
    )
    if declared is None and not observed:
        raise ValueError(
            f"Rules row {rule.row_number}: field path {rule.field_path!r} is not declared or present"
        )
    if declared is not None and declared.inferred:
        raise ValueError(f"Rules row {rule.row_number}: inferred fields cannot be transformed")
    return file_model


def _samples_before(
    file_model: DataFile, field_path: str
) -> tuple[tuple[int, str, str], ...]:
    """Capture scalar target values for a bounded before/after response.

    Require: field path was resolved for the current File.
    Guarantee: return (one-based ordinal, Entry name, scalar value) rows.
    Maintain: do not mutate Entries or infer values for missing fields.
    Assert: each sampled value is source text.
    """
    result: list[tuple[int, str, str]] = []
    for ordinal, entry in enumerate(file_model.entries, start=1):
        instance = file_model._mass_modify_instance(entry, field_path)
        if instance is not None and isinstance(instance.value, str):
            if not instance.field.inferred and not (
                instance.field.arguments
                and instance.field.arguments[0].argument_type.startswith("*")
            ):
                result.append((ordinal, entry.name, instance.value))
    return tuple(result)


def _execute_rules(
    document: LoadedProject, rules: tuple[TransformRule, ...]
) -> TransformRun:
    """Execute ordered rules on the request's disposable project model.

    Require: every rule passed workbook validation and document is freshly loaded.
    Guarantee: return changed-entry samples and changed bytes without disk writes.
    Maintain: rules are sequential; each formula is internally source-snapshotted.
    Assert: each changed File serializes before it can be published.
    """
    results: list[RuleResult] = []
    for rule in rules:
        file_model = _rule_file(document, rule)
        before = _samples_before(file_model, rule.field_path)
        reference = lambda structure_file, entry_name, field_path: (
            resolve_project_formula_reference(
                document, structure_file, entry_name, field_path
            )
        )
        with tqdm(
            total=file_model.entries.count,
            desc=f"{rule.structure_file}:{rule.field_path}",
            unit="entry",
            leave=False,
        ) as progress:
            def update_progress(completed: int, total: int) -> None:
                """Update the terminal bar from one completed model traversal.

                Require: ``completed`` is within the supplied batch total.
                Guarantee: advance the bar to the completed Entry count.
                Maintain: progress affects only terminal output.
                Assert: model and tqdm totals agree.
                """
                if total != progress.total:
                    raise AssertionError("progress total differs from the File Entry count")
                progress.update(completed - progress.n)

            changed_count = file_model.mass_modify(
                rule.field_path,
                rule.operation,
                rule.operand,
                reference_resolver=reference if rule.operation == "formula" else None,
                progress_callback=update_progress,
            )
        after = _samples_before(file_model, rule.field_path)
        after_by_ordinal = {ordinal: (name, value) for ordinal, name, value in after}
        samples = tuple(
            {
                "ordinal": ordinal,
                "entry": after_by_ordinal[ordinal][0],
                "before": old_value,
                "after": after_by_ordinal[ordinal][1],
            }
            for ordinal, _name, old_value in before
            if ordinal in after_by_ordinal and old_value != after_by_ordinal[ordinal][1]
        )[:10]
        results.append(
            RuleResult(
                rule.row_number,
                rule.structure_file,
                rule.field_path,
                rule.operation,
                changed_count,
                samples,
            )
        )

    changed_files: list[tuple[Path, bytes]] = []
    data_root = document.data_root.resolve()
    for file_model in document.project.files:
        source = Path(file_model.source).resolve()
        try:
            relative = source.relative_to(data_root)
        except ValueError as error:
            if file_model.to_bytes() != file_model.original_bytes:
                raise ValueError(
                    f"changed File {source} is outside the project's data root"
                ) from error
            continue
        serialized = file_model.to_bytes()
        if serialized != file_model.original_bytes:
            if relative.is_absolute() or ".." in relative.parts:
                raise ValueError(f"changed File path escapes the project data root: {source}")
            changed_files.append((relative, serialized))
    return TransformRun(tuple(results), tuple(changed_files), data_root)


def _run_upload(
    workbook_data: bytes,
    gamefile_path: str,
    template_path: str,
    structure_root: str,
    language: str | None,
) -> TransformRun:
    """Parse and execute one preview/apply request without file writes.

    Require: workbook bytes and explicit project paths are supplied.
    Guarantee: return a completed in-memory transformation run.
    Maintain: each request reloads the source project from disk.
    Assert: source Files remain byte-identical after simulation.
    """
    rules = parse_rules_workbook(workbook_data)
    document = _load_document(gamefile_path, template_path, structure_root, language)
    return _execute_rules(document, rules)


def _response_payload(run: TransformRun) -> dict[str, Any]:
    """Convert an in-memory result to the stable JSON contract.

    Require: run contains validated relative Paths and JSON-safe samples.
    Guarantee: return ordered rules and POSIX relative file names.
    Maintain: do not disclose local source paths.
    Assert: changed paths are nonabsolute.
    """
    return {
        "rules": [asdict(result) for result in run.results],
        "changed_files": [path.as_posix() for path, _data in run.changed_files],
    }


def _publish_overlay(
    run: TransformRun, output_directory: str
) -> Path:
    """Publish a complete changed-file overlay to a new destination.

    Require: output parent exists; output does not exist; changed paths are safe.
    Guarantee: destination contains all changed files or no destination at all.
    Maintain: source data root is never a destination or ancestor of output.
    Assert: each file is staged beneath a temporary sibling before rename.
    """
    destination = Path(output_directory).resolve()
    source_root = run.data_root.resolve()
    if destination.exists():
        raise ValueError("output_directory must not already exist")
    if destination == source_root or destination.is_relative_to(source_root):
        raise ValueError("output_directory must be outside the source data root")
    if not destination.parent.is_dir():
        raise ValueError("output_directory parent must already exist")

    temporary = Path(
        tempfile.mkdtemp(prefix=f".{destination.name}.staging-", dir=destination.parent)
    )
    try:
        for relative, data in run.changed_files:
            if relative.is_absolute() or ".." in relative.parts:
                raise ValueError(f"output path escapes overlay: {relative}")
            target = temporary.joinpath(*relative.parts)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)
        os.rename(temporary, destination)
    finally:
        if temporary.exists():
            shutil.rmtree(temporary)
    return destination


app = FastAPI(title="CTPEd transformation API", version="1.0.0")


@app.get("/health")
def health() -> dict[str, str]:
    """Report local readiness without touching project files.

    Require: FastAPI has dispatched GET /health.
    Guarantee: return the stable ``status=ok`` response.
    Maintain: do not inspect local paths or project state.
    Assert: response values are JSON strings.
    """
    return {"status": "ok"}


@app.get("/v1/transforms/template")
def workbook_template() -> Response:
    """Return the blank workbook with the supported rule headers.

    Require: openpyxl is installed through the optional API dependencies.
    Guarantee: return XLSX bytes with a download filename and content type.
    Maintain: generate the workbook in memory only.
    Assert: workbook rows are not run as transformations.
    """
    return Response(
        create_rules_workbook(),
        media_type=WORKBOOK_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="ctpedit-transform-rules.xlsx"'},
    )


def _read_upload(workbook: UploadFile) -> bytes:
    """Read one bounded XLSX upload.

    Require: ``workbook`` is the UploadFile received by a route.
    Guarantee: return at most 10 MiB plus one detection byte.
    Maintain: do not persist uploads beyond FastAPI's request spool.
    Assert: the filename has the supported .xlsx suffix.
    """
    if not workbook.filename or Path(workbook.filename).suffix.casefold() != ".xlsx":
        raise ValueError("workbook upload must use the .xlsx extension")
    data = workbook.file.read(MAX_WORKBOOK_BYTES + 1)
    if len(data) > MAX_WORKBOOK_BYTES:
        raise ValueError("workbook upload exceeds the 10 MiB limit")
    return data


def _http_error(error: Exception) -> HTTPException:
    """Map caller-controlled input and model failures to HTTP 422.

    Require: ``error`` describes invalid request data or an unusable project.
    Guarantee: preserve its actionable message in response detail.
    Maintain: do not expose a traceback or server environment.
    Assert: the returned exception uses the client-error status.
    """
    return HTTPException(status_code=422, detail=str(error))


@app.post("/v1/transforms/preview")
def preview_transform(
    workbook: UploadFile = UploadField(...),
    gamefile_path: str = Form(...),
    template_path: str = Form(...),
    structure_root: str = Form(...),
    language: str | None = Form(default=None),
) -> dict[str, Any]:
    """Preview the uploaded rule sequence without writing files.

    Require: workbook and project path fields are supplied as multipart data.
    Guarantee: return ordered counts and samples or HTTP 422.
    Maintain: only a freshly loaded in-memory project can be changed.
    Assert: source data bytes remain untouched.
    """
    try:
        run = _run_upload(
            _read_upload(workbook),
            gamefile_path,
            template_path,
            structure_root,
            language,
        )
    except (OSError, ValueError, TypeError, IndexError, ArithmeticError, NotImplementedError) as error:
        raise _http_error(error) from error
    finally:
        workbook.file.close()
    return _response_payload(run)


@app.post("/v1/transforms/apply")
def apply_transform(
    workbook: UploadFile = UploadField(...),
    gamefile_path: str = Form(...),
    template_path: str = Form(...),
    structure_root: str = Form(...),
    output_directory: str = Form(...),
    language: str | None = Form(default=None),
) -> dict[str, Any]:
    """Apply rules to a fresh project and publish a new overlay.

    Require: multipart fields include a new output directory outside data_root.
    Guarantee: return rule results and overlay path or HTTP 422.
    Maintain: never write to source Gamefile, descriptor, or data paths.
    Assert: errors leave no published partial overlay.
    """
    try:
        workbook_data = _read_upload(workbook)
        run = _run_upload(
            workbook_data,
            gamefile_path,
            template_path,
            structure_root,
            language,
        )
        destination = _publish_overlay(run, output_directory)
    except (OSError, ValueError, TypeError, IndexError, ArithmeticError, NotImplementedError) as error:
        raise _http_error(error) from error
    finally:
        workbook.file.close()
    result = _response_payload(run)
    result["output_directory"] = str(destination)
    return result


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="127.0.0.1", port=8765)
